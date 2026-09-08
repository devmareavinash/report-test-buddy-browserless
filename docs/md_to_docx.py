from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Inches


def add_runs_with_code(paragraph, text):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if part.startswith("`") and part.endswith("`") and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            sub = re.split(r"(\*\*[^*]+\*\*)", part)
            for s in sub:
                if s.startswith("**") and s.endswith("**") and len(s) >= 4:
                    run = paragraph.add_run(s[2:-2])
                    run.bold = True
                else:
                    paragraph.add_run(s)


def md_to_docx(md_path: Path, docx_path: Path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = []
        for row in table_buf:
            if re.match(r"^\|?\s*-+", row):
                continue
            cells = [c.strip() for c in row.strip().strip("|").split("|")]
            rows.append(cells)
        table_buf = []
        if not rows:
            return
        cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for r_idx, row in enumerate(rows):
            for c_idx in range(cols):
                val = row[c_idx] if c_idx < len(row) else ""
                cell = table.rows[r_idx].cells[c_idx]
                cell.text = ""
                p = cell.paragraphs[0]
                add_runs_with_code(p, val)
                if r_idx == 0:
                    for run in p.runs:
                        run.bold = True
        doc.add_paragraph()

    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            if i >= len(lines) or not lines[i].strip().startswith("|"):
                flush_table()
            continue

        if table_buf:
            flush_table()

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
        elif line.strip() == "---":
            doc.add_paragraph("\u2500" * 40)
        elif re.match(r"^[-*] ", line.strip()):
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_code(p, line.strip()[2:])
        elif re.match(r"^\d+\. ", line.strip()):
            p = doc.add_paragraph(style="List Number")
            add_runs_with_code(p, re.sub(r"^\d+\.\s+", "", line.strip()))
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph()
            add_runs_with_code(p, line)

        i += 1

    flush_table()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"Wrote {docx_path}")


if __name__ == "__main__":
    base = Path(
        r"c:\Users\EASXP\OneDrive - Bayer\Desktop\Report Test Buddy Brwserless\docs"
    )
    md_to_docx(
        base / "ARCHITECTURE_AND_FUNCTIONING.md",
        base / "ARCHITECTURE_AND_FUNCTIONING.docx",
    )
    md_to_docx(base / "DEPLOYMENT_GUIDE.md", base / "DEPLOYMENT_GUIDE.docx")
    print("Done")
